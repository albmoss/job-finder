"""
Wyciąganie tekstu z CV w formacie PDF i DOCX.
"""

import logging
from pathlib import Path
from typing import Optional

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

logger = logging.getLogger(__name__)


class CVParser:
    """Wyciąga tekst z pliku CV."""
    
    @staticmethod
    def extract_from_pdf(filepath: str) -> Optional[str]:
        if PdfReader is None:
            raise ImportError("pypdf not installed. Install with: pip install pypdf")
        
        try:
            reader = PdfReader(filepath)
            text_parts = []
            
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            
            full_text = "\n".join(text_parts)
            logger.info(f"Extracted {len(full_text)} characters from PDF: {filepath}")
            return full_text
            
        except Exception as e:
            logger.error(f"Error extracting PDF {filepath}: {e}")
            return None
    
    @staticmethod
    def extract_from_docx(filepath: str) -> Optional[str]:
        if Document is None:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")
        
        try:
            doc = Document(filepath)
            text_parts = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Tekst z tabel też
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            full_text = "\n".join(text_parts)
            logger.info(f"Extracted {len(full_text)} characters from DOCX: {filepath}")
            return full_text
            
        except Exception as e:
            logger.error(f"Error extracting DOCX {filepath}: {e}")
            return None
    
    @staticmethod
    def parse_cv(filepath: str) -> Optional[str]:
        """
        Auto-detect file type and extract text
        
        Args:
            filepath: Path to CV file (PDF or DOCX)
            
        Returns:
            Extracted text or None if extraction failed
        """
        path = Path(filepath)
        
        if not path.exists():
            logger.error(f"File not found: {filepath}")
            return None
        
        extension = path.suffix.lower()
        
        if extension == '.pdf':
            return CVParser.extract_from_pdf(filepath)
        elif extension in ['.docx', '.doc']:
            return CVParser.extract_from_docx(filepath)
        else:
            logger.error(f"Unsupported file format: {extension}")
            return None
    
    @staticmethod
    def clean_text(text: str) -> str:
        if not text:
            return ""
        
        # Zbijamy nadmiarowe białe znaki
        lines = [line.strip() for line in text.split('\n')]
        lines = [line for line in lines if line]
        
        cleaned = "\n".join(lines)
        return cleaned
